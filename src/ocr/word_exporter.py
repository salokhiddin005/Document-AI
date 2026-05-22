"""
Word (.docx) Exporter — converts OCR text into a formatted Word document.
"""

from __future__ import annotations

from pathlib import Path
from loguru import logger


def export_to_docx(
    text: str,
    output_path: Path,
    title: str = "",
    lang: str = "en",
) -> Path:
    """
    Export OCR text as a formatted .docx file.
    Preserves paragraph breaks. Adds title if provided.
    """
    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")

    doc = Document()

    # Page margins (A4 standard)
    for section in doc.sections:
        section.page_width  = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin  = section.bottom_margin = Cm(2.5)

    # Title
    if title:
        t = doc.add_heading(title, level=1)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Body — split on double newlines for paragraphs
    paragraphs = text.split("\n\n")
    for para_text in paragraphs:
        para_text = para_text.replace("\n", " ").strip()
        if not para_text:
            continue
        p = doc.add_paragraph(para_text)
        p.paragraph_format.space_after = Pt(6)
        for run in p.runs:
            run.font.size = Pt(12)
            run.font.name = "Times New Roman"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    logger.info(f"Word document saved: {output_path}")
    return output_path
